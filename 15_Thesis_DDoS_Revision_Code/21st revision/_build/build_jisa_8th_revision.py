#!/usr/bin/env python3
"""Build the JISA manuscript from the executed notebook counts.

Corrects the protocol and replaces rounded rates with the printed confusion
counts. Does not invent matched-threshold, overlap, or repeated-seed results.
"""
from __future__ import annotations

import csv
import json
import re
import shutil
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_LINE_SPACING
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font

ROOT = Path(r"D:\Research\Articles\MSc research\20th revision")
RIS_ROOT = Path(r"D:\Research\Articles\MSc research\journal articles\RIS Files")
sys.path.insert(0, str(ROOT / "_build"))

import derived_metrics as dm  # noqa: E402
from reproduce_derived_analyses import write_all  # noqa: E402

OUT = ROOT / "Manuscript uploads"
OUT_FIG = OUT / "Figures"
OUT_TAB = OUT / "Tables"
TABLES = ROOT / "02_Tables"
ZENODO = ROOT / "zenodo uploads"
SUB = ROOT / "01_Submission_Files"

TITLE = (
    "High F1, Costly Alerts: Dataset- and Prevalence-Dependent Rankings of "
    "Deep and Tree-Based DDoS Detectors"
)
NUMS = dm.key_numbers()
H = dm.by_name(dm.HYBRID)
RF = dm.by_name(dm.RF)
ATTACKS_AT_1PCT = 10_000_000 * 0.01 / 0.99
MISS_H = H["FNR"] / 100.0 * ATTACKS_AT_1PCT
MISS_RF = RF["FNR"] / 100.0 * ATTACKS_AT_1PCT

ABSTRACT = (
    "Public DDoS benchmarks often rank detectors by F1. A CNN–LSTM hybrid and a random forest both exceed 99% F1, "
    "but at the evaluated points their false-alarm order changes with the collection. On CICDDoS2019, counts "
    "recovered from the printed precision, recall, support, and Matthews correlation give the forest no false "
    "alarms in 10,663 benign records and the hybrid 19 in 10,650 benign windows. The forest also misses fewer "
    "attacks. The binomial intervals on those rates do not overlap only if each window is an independent trial. "
    "On CICIoT2023 the order reverses at one cut on the scored slices: one hybrid false alarm in 5,962 windows "
    "against 431 forest false alarms in 7,064 records. Holding the CICIoT rates fixed, projected precision at "
    f"1% attack prevalence is {NUMS['hybrid_precision_1pct']:.1f}% and {NUMS['rf_precision_1pct']:.1f}%, and equal "
    f"assumed costs switch the cheaper detector at {NUMS['crossover_equal_pct']:.1f}% prevalence. At the CICIoT2023 "
    "training mix, when a miss is assumed to cost 100 times a false alarm, always alerting can cost less than "
    "either learned detector. On CICDDoS2019 a decision tree places about 99% of its importance on an Inbound "
    "flag and two port fields. On a four-file subset, removing that flag and the ports raises a forest from 1 to "
    "5 false alarms in 3,011 benign records. On a BoT-IoT mirror subsample that keeps all 477 benign rows, the "
    "seed-42 tuned cut leaves a decision tree at 175 misses. Rankings should be read with the dataset, the scored "
    "cases, the threshold, and the assumed prevalence and costs."
)

HIGHLIGHTS = [
    "At the evaluated points, the false-alarm order flips",
    "CICDDoS2019: forest has 0 false alarms; the hybrid has 19",
    f"At 1% prevalence, projected precision is {NUMS['hybrid_precision_1pct']:.1f}% vs {NUMS['rf_precision_1pct']:.1f}%",
    "A tree puts 99% of importance on an Inbound flag and two ports",
    "BoT-IoT subsample: the seed-42 tuned cut stays at 175 misses",
]

KEYWORDS = [
    "DDoS detection",
    "IoT security",
    "false alarm rate",
    "attack prevalence",
    "cost-sensitive evaluation",
    "shortcut learning",
    "base-rate fallacy",
]

AFF = {
    "a": "Department of Electrical and Computer Engineering (DEEC), University of Coimbra, Rua Sílvio Lima, Polo II, 3030-290 Coimbra, Portugal",
    "b": "Department of Mechanics, Faculty of Civil Engineering, Czech Technical University in Prague, Thákurova 7, 166 29 Prague 6, Czech Republic",
}
AUTHORS = [
    ("Hina Iqbal", "a", "hinazoha4@gmail.com", "0009-0002-2378-9414", True),
    ("Cristiano Premebida", "a", "cpremebida@deec.uc.pt", "0000-0002-2168-2077", False),
    ("Usman Haider", "b", "usmanhaider6886@gmail.com", "0000-0002-0125-0912", False),
]


def set_run_font(run, size=11, bold=False, italic=False, name="Times New Roman"):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    run.font.size = Pt(size)
    run.bold = bold
    run.italic = italic
    run.font.superscript = False


def add_para(doc, text="", size=11, bold=False, italic=False, align="left", space_after=8, first_line=False, double=True):
    p = doc.add_paragraph()
    if align == "center":
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    elif align == "justify":
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    if double:
        p.paragraph_format.line_spacing_rule = WD_LINE_SPACING.DOUBLE
    p.paragraph_format.space_after = Pt(space_after)
    if first_line:
        p.paragraph_format.first_line_indent = Inches(0.5)
    if text:
        run = p.add_run(text)
        set_run_font(run, size=size, bold=bold, italic=italic)
    return p


def add_h(doc, text, level=1):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6)
    run = p.add_run(text)
    set_run_font(run, size=12 if level == 1 else 11, bold=True)
    return p


def add_rich(doc, parts, size=11, align="justify", first_line=False, space_after=8):
    p = add_para(doc, align=align, first_line=first_line, space_after=space_after)
    for text, italic in parts:
        run = p.add_run(text)
        set_run_font(run, size=size, italic=italic)
    return p


def horizontal_borders(table):
    tbl_pr = table._tbl.tblPr
    borders = OxmlElement("w:tblBorders")
    for edge, val, sz in (
        ("top", "single", "8"),
        ("left", "nil", "0"),
        ("bottom", "single", "8"),
        ("right", "nil", "0"),
        ("insideH", "single", "4"),
        ("insideV", "nil", "0"),
    ):
        element = OxmlElement(f"w:{edge}")
        element.set(qn("w:val"), val)
        element.set(qn("w:sz"), sz)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), "000000")
        borders.append(element)
    tbl_pr.append(borders)


def load_csv(name: str) -> list[dict]:
    with (TABLES / name).open(encoding="utf-8") as handle:
        return list(csv.DictReader(handle))


def embed_table(doc, rows, caption, note, columns=None):
    add_para(doc, caption, size=10, italic=True, space_after=4, double=False)
    keys = columns or list(rows[0].keys())
    table = doc.add_table(rows=1 + len(rows), cols=len(keys))
    table.style = "Table Grid"
    horizontal_borders(table)
    for i, key in enumerate(keys):
        table.rows[0].cells[i].text = key.replace("_", " ")
        for paragraph in table.rows[0].cells[i].paragraphs:
            for run in paragraph.runs:
                set_run_font(run, size=8, bold=True)
    for r_i, row in enumerate(rows):
        for c_i, key in enumerate(keys):
            table.rows[r_i + 1].cells[c_i].text = str(row.get(key, ""))
            for paragraph in table.rows[r_i + 1].cells[c_i].paragraphs:
                for run in paragraph.runs:
                    set_run_font(run, size=8)
    add_para(doc, note, size=9, italic=True, space_after=8, double=False)


def add_procedure(doc):
    add_para(doc, "Box 1. Prevalence-aware reading of one operating point.", bold=True, space_after=4)
    steps = [
        "1. Fix the benign volume and the largest false-alert load that can be handled. Their ratio is the false-alarm budget.",
        "2. On validation data only, choose a threshold for each detector that meets that budget, then compare miss rates on the same test cases. Section 5.5 does this on BoT-IoT. The CIC results in Section 5 use the argmax cut on the slices in Table 6, which are not aligned.",
        "3. Keep the points on the ROC convex hull of the available cuts. For the hybrid and random forest at the evaluated cut, equation (3) gives the prevalence at which their expected-cost order changes. Apply the never-alert and always-alert bounds before treating either model as the cheapest policy.",
        "4. Report projected precision and false alerts at the deployment prevalence, and state the cost ratio as a scenario rather than as a measured price. State which findings were counted and which were projected.",
    ]
    for step in steps:
        add_para(doc, step, align="justify", space_after=4)


def csv_to_xlsx(path: Path):
    with path.open(encoding="utf-8") as handle:
        rows = list(csv.reader(handle))
    wb = Workbook()
    ws = wb.active
    ws.title = "data"
    for row in rows:
        ws.append(row)
    for cell in ws[1]:
        cell.font = Font(bold=True, name="Times New Roman")
        cell.alignment = Alignment(wrap_text=True)
    wb.save(path.with_suffix(".xlsx"))


def refs() -> list[str]:
    return [
        "Akiba, T., Sano, S., Yanase, T., Ohta, T., & Koyama, M. (2019). Optuna: A next-generation hyperparameter optimization framework. In Proceedings of the 25th ACM SIGKDD International Conference on Knowledge Discovery & Data Mining (pp. 2623–2631). https://doi.org/10.1145/3292500.3330701",
        "Anurupam, K., & Seshadri, K. (2025). Lightweight representation learning for network traffic towards malicious traffic detection in edge devices. Journal of Information Security and Applications, 93, 104186. https://doi.org/10.1016/j.jisa.2025.104186",
        "Apruzzese, G., Laskov, P., & Schneider, J. (2023). SoK: Pragmatic assessment of machine learning for network intrusion detection. In 2023 IEEE European Symposium on Security and Privacy (EuroS&P) (pp. 592–614). https://doi.org/10.1109/EuroSP57164.2023.00042",
        "Arp, D., Quiring, E., Pendlebury, F., Warnecke, A., Pierazzi, F., Wressnegger, C., Cavallaro, L., & Rieck, K. (2022). Dos and don’ts of machine learning in computer security. In 31st USENIX Security Symposium. https://www.usenix.org/conference/usenixsecurity22/presentation/arp",
        "Axelsson, S. (2000). The base-rate fallacy and the difficulty of intrusion detection. ACM Transactions on Information and System Security, 3(3), 186–205. https://doi.org/10.1145/357830.357849",
        "Bergstra, J., Bardenet, R., Bengio, Y., & Kégl, B. (2011). Algorithms for hyper-parameter optimization. In Advances in Neural Information Processing Systems 24 (pp. 2546–2554).",
        "Bizzarri, A., Yu, C.-E., Jalaian, B., Riguzzi, F., & Bastian, N. D. (2025). Neurosymbolic AI for network intrusion detection systems: A survey. Journal of Information Security and Applications, 94, 104205. https://doi.org/10.1016/j.jisa.2025.104205",
        "Canadian Institute for Cybersecurity. (2019). CICDDoS2019 [Dataset]. University of New Brunswick. https://www.unb.ca/cic/datasets/ddos-2019.html (accessed 15 October 2026).",
        "Canadian Institute for Cybersecurity. (2023). CICIoT2023 [Dataset]. University of New Brunswick. https://www.unb.ca/cic/datasets/iotdataset-2023.html (accessed 15 October 2026).",
        "Chicco, D., & Jurman, G. (2020). The advantages of the Matthews correlation coefficient (MCC) over F1 score and accuracy in binary classification evaluation. BMC Genomics, 21, 6. https://doi.org/10.1186/s12864-019-6413-7",
        "Cogranne, R., Letourneau, M., & Doyen, G. (2026). Deep learning-based sequential detection of attacks on low-latency network services. Journal of Information Security and Applications, 100, 104458. https://doi.org/10.1016/j.jisa.2026.104458",
        "Coscia, A., Dentamaro, V., Galantucci, S., Maci, A., & Pirlo, G. (2024). Automatic decision tree-based NIDPS ruleset generation for DoS/DDoS attacks. Journal of Information Security and Applications, 82, 103736. https://doi.org/10.1016/j.jisa.2024.103736",
        "D’hooge, L., Wauters, T., Volckaert, B., & De Turck, F. (2020). Inter-dataset generalization strength of supervised machine learning methods for intrusion detection. Journal of Information Security and Applications, 54, 102564. https://doi.org/10.1016/j.jisa.2020.102564",
        "Doriguzzi-Corin, R., Millar, S., Scott-Hayward, S., Martinez-del-Rincon, J., & Siracusa, D. (2020). LUCID: A practical, lightweight deep learning solution for DDoS attack detection. IEEE Transactions on Network and Service Management, 17(2), 876–889. https://doi.org/10.1109/TNSM.2020.2971776",
        "Elkan, C. (2001). The foundations of cost-sensitive learning. In Proceedings of the 17th International Joint Conference on Artificial Intelligence (pp. 973–978).",
        "ElShehaby, M., & Matrawy, A. (2026). A novel perturb-ability score to mitigate evasion adversarial attacks on flow-based ML-NIDS. Journal of Information Security and Applications, 99, 104409. https://doi.org/10.1016/j.jisa.2026.104409",
        "Engelen, G., Rimmer, V., & Joosen, W. (2021). Troubleshooting an intrusion detection dataset: The CICIDS2017 case study. In 2021 IEEE Security and Privacy Workshops (SPW) (pp. 7–12). https://doi.org/10.1109/SPW53761.2021.00009",
        "Fawcett, T. (2006). An introduction to ROC analysis. Pattern Recognition Letters, 27(8), 861–874. https://doi.org/10.1016/j.patrec.2005.10.010",
        "Ferrag, M. A., Maglaras, L., Moschoyiannis, S., & Janicke, H. (2020). Deep learning for cyber security intrusion detection: Approaches, datasets, and comparative study. Journal of Information Security and Applications, 50, 102419. https://doi.org/10.1016/j.jisa.2019.102419",
        "Flood, R., Engelen, G., Aspinall, D., & Desmet, L. (2024). Bad design smells in benchmark NIDS datasets. In 2024 IEEE European Symposium on Security and Privacy (EuroS&P). https://doi.org/10.1109/EuroSP60621.2024.00042",
        "González-Ramos, J. A., Abril, E. J., Fernández, P., Prieto, J., & Chamoso, P. (2026). Adversarial robustness evaluation of hybrid CNN-LSTM-transformer NIDS on evolving threats. Journal of Information Security and Applications, 100, 104467. https://doi.org/10.1016/j.jisa.2026.104467",
        "Grinsztajn, L., Oyallon, E., & Varoquaux, G. (2022). Why do tree-based models still outperform deep learning on typical tabular data? In Advances in Neural Information Processing Systems (Datasets and Benchmarks). https://openreview.net/forum?id=Ob2IdlY7sA",
        "Gu, G., Fogla, P., Dagon, D., Lee, W., & Škorić, B. (2006). Measuring intrusion detection capability: An information-theoretic approach. In Proceedings of the ACM Symposium on Information, Computer and Communications Security (pp. 90–101). https://doi.org/10.1145/1128817.1128834",
        "Guo, C., Pleiss, G., Sun, Y., & Weinberger, K. Q. (2017). On calibration of modern neural networks. In Proceedings of the 34th International Conference on Machine Learning (pp. 1321–1330). PMLR. https://proceedings.mlr.press/v70/guo17a.html",
        "Hochreiter, S., & Schmidhuber, J. (1997). Long short-term memory. Neural Computation, 9(8), 1735–1780. https://doi.org/10.1162/neco.1997.9.8.1735",
        "Hormozi, M., Erfani, S. H., Sahafi, A., & Moradi, M. (2026). Fed-Adapt: A federated learning framework for adaptive topology reconfiguration against multi-rate DDoS and database flooding attacks. Journal of Information Security and Applications, 98, 104384. https://doi.org/10.1016/j.jisa.2026.104384",
        "Iqbal, H. (2026). Bayesian-optimized hybrid deep learning framework for detecting DDoS attacks on MEC and IoT networks (Master’s dissertation). University of Coimbra. Available from the corresponding author on request.",
        "Iqbal, H., Premebida, C., & Haider, U. (2026). Supporting research materials for “High F1, Costly Alerts: Dataset- and Prevalence-Dependent Rankings of Deep and Tree-Based DDoS Detectors” (interim version) [Data set]. Zenodo. https://doi.org/10.5281/zenodo.23435601",
        "Javadpour, A., Ja’fari, F., Benzaïd, C., & Taleb, T. (2025). An optimized reinforcement learning based MTD mutation strategy for securing edge IoT against DDoS attack. Journal of Information Security and Applications, 93, 104138. https://doi.org/10.1016/j.jisa.2025.104138",
        "Kapoor, S., & Narayanan, A. (2023). Leakage and the reproducibility crisis in machine-learning-based science. Patterns, 4(9), 100804. https://doi.org/10.1016/j.patter.2023.100804",
        "Ko, I., Chambers, D., & Barrett, E. (2020). Adaptable feature-selecting and threshold-moving complete autoencoder for DDoS flood attack mitigation. Journal of Information Security and Applications, 55, 102647. https://doi.org/10.1016/j.jisa.2020.102647",
        "Koroniotis, N., Moustafa, N., Sitnikova, E., & Turnbull, B. (2019). Towards the development of realistic botnet dataset in the internet of things for network forensic analytics: Bot-IoT dataset. Future Generation Computer Systems, 100, 779–796. https://doi.org/10.1016/j.future.2019.02.050",
        "Lara-Gutierrez, A., Fernandez-Gago, C., & Onieva, J. A. (2026). Adversarial red teaming and imbalance correction in heterogeneous NIDS: A class-specific adaptive GAN framework. Journal of Information Security and Applications, 102, 104584. https://doi.org/10.1016/j.jisa.2026.104584",
        "Layeghy, S., Gallagher, M., & Portmann, M. (2024). Benchmarking the benchmark — Comparing synthetic and real-world Network IDS datasets. Journal of Information Security and Applications, 80, 103689. https://doi.org/10.1016/j.jisa.2023.103689",
        "Li, M., Zheng, L., Ma, X., & Li, S. (2025). Real-time monitoring model of DDoS attacks using distance thresholds in Edge cooperation networks. Journal of Information Security and Applications, 89, 103972. https://doi.org/10.1016/j.jisa.2025.103972",
        "Madoune, S. A., Senouci, S., Jiang, D. D., Senouci, M. R., Daoud, M. A., Anwar Mohammed Alawad, R., et al. (2025). A novel approach for real-time DDoS detection in SDN using dimensionality reduction and ensemble learning. Journal of Information Security and Applications, 94, 104195. https://doi.org/10.1016/j.jisa.2025.104195",
        "Meidan, Y., Bohadana, M., Mathov, Y., Mirsky, Y., Shabtai, A., Breitenbacher, D., & Elovici, Y. (2018). N-BaIoT: Network-based detection of IoT botnet attacks using deep autoencoders. IEEE Pervasive Computing, 17(3), 12–22. https://doi.org/10.1109/MPRV.2018.03367731",
        "Mittal, M., Kumar, K., & Behal, S. (2023). DL-2P-DDoSADF: Deep learning-based two-phase DDoS attack detection framework. Journal of Information Security and Applications, 78, 103609. https://doi.org/10.1016/j.jisa.2023.103609",
        "Neto, E. C. P., Dadkhah, S., Ferreira, R., Zohourian, A., Lu, R., & Ghorbani, A. A. (2023). CICIoT2023: A real-time dataset and benchmark for large-scale attacks in IoT environment. Sensors, 23(13), 5941. https://doi.org/10.3390/s23135941",
        "Pedregosa, F., Varoquaux, G., Gramfort, A., Michel, V., Thirion, B., Grisel, O., et al. (2011). Scikit-learn: Machine learning in Python. Journal of Machine Learning Research, 12, 2825–2830.",
        "Pendlebury, F., Pierazzi, F., Jordaney, R., Kinder, J., & Cavallaro, L. (2019). TESSERACT: Eliminating experimental bias in malware classification across space and time. In Proceedings of the 28th USENIX Security Symposium. https://www.usenix.org/conference/usenixsecurity19/presentation/pendlebury",
        "Provost, F., & Fawcett, T. (2001). Robust classification for imprecise environments. Machine Learning, 42(3), 203–231. https://doi.org/10.1023/A:1007601015854",
        "Ring, M., Wunderlich, S., Scheuring, D., Landes, D., & Hotho, A. (2019). A survey of network-based intrusion detection data sets. Computers & Security, 86, 147–167. https://doi.org/10.1016/j.cose.2019.06.005",
        "Rustam, F., & Jurcut, A. D. (2025). Adaptive security framework for multi-environment networks using ensemble data drift detection and incremental deep learning. Journal of Information Security and Applications, 94, 104219. https://doi.org/10.1016/j.jisa.2025.104219",
        "Saito, T., & Rehmsmeier, M. (2015). The precision-recall plot is more informative than the ROC plot when evaluating binary classifiers on imbalanced datasets. PLoS ONE, 10(3), e0118432. https://doi.org/10.1371/journal.pone.0118432",
        "Sharafaldin, I., Lashkari, A. H., Hakak, S., & Ghorbani, A. A. (2019). Developing realistic distributed denial of service (DDoS) attack dataset and taxonomy. In 2019 International Carnahan Conference on Security Technology (ICCST) (pp. 1–8). https://doi.org/10.1109/CCST.2019.8888419",
        "Shi, W., Cao, J., Zhang, Q., Li, Y., & Xu, L. (2016). Edge computing: Vision and challenges. IEEE Internet of Things Journal, 3(5), 637–646. https://doi.org/10.1109/JIOT.2016.2579198",
        "Shwartz-Ziv, R., & Armon, A. (2022). Tabular data: Deep learning is not all you need. Information Fusion, 81, 84–90. https://doi.org/10.1016/j.inffus.2021.11.011",
        "Singha, M. F., Patgiri, R., & Singh, L. D. (2025). A tenant-aware deep learning-based intrusion detection system for detecting DDoS attacks in multi-tenant SaaS networks. Journal of Information Security and Applications, 94, 104251. https://doi.org/10.1016/j.jisa.2025.104251",
        "Snoek, J., Larochelle, H., & Adams, R. P. (2012). Practical Bayesian optimization of machine learning algorithms. In Advances in Neural Information Processing Systems 25 (pp. 2951–2959).",
        "Sommer, R., & Paxson, V. (2010). Outside the closed world: On using machine learning for network intrusion detection. In IEEE Symposium on Security and Privacy (pp. 305–316). https://doi.org/10.1109/SP.2010.25",
        "Vaarandi, R., & Guerra-Manzanares, A. (2024). Network IDS alert classification with active learning techniques. Journal of Information Security and Applications, 81, 103687. https://doi.org/10.1016/j.jisa.2023.103687",
        "Wang, S., Zhao, Y., Fu, X., Si, H., Wang, W., & Xue, L. (2026a). A class-imbalance-aware intrusion detection system based on spatiotemporal graph neural networks for software-defined vehicles. Journal of Information Security and Applications, 97, 104340. https://doi.org/10.1016/j.jisa.2025.104340",
        "Wang, Y., Yan, Z., Zhang, K., & Wen, M. (2026b). DEGAN: Towards botnet detection in IIoT with dual-enhanced GAN under imbalanced data. Journal of Information Security and Applications, 100, 104499. https://doi.org/10.1016/j.jisa.2026.104499",
        "Wolf, M., Bergner, K., Förtsch, P., & Landes, D. (2026). Implications of applying emulation data to machine learning-based intrusion detection systems. Journal of Information Security and Applications, 97, 104362. https://doi.org/10.1016/j.jisa.2025.104362",
        "Zargar, S. T., Joshi, J., & Tipper, D. (2013). A survey of defense mechanisms against distributed denial of service (DDoS) flooding attacks. IEEE Communications Surveys & Tutorials, 15(4), 2046–2069. https://doi.org/10.1109/SURV.2013.031413.00127",
        "Zhao, X., Fok, K. W., & Thing, V. L. L. (2026). Threshold-free network anomaly detection via comparative reconstruction error learning with parallel GANs. Journal of Information Security and Applications, 101, 104536. https://doi.org/10.1016/j.jisa.2026.104536",
    ]


def screen_ris() -> dict:
    records = []
    for path in RIS_ROOT.rglob("*.ris"):
        text = path.read_text(encoding="utf-8", errors="replace")
        for block in re.split(r"\nER\s+-\s*\n", text):
            if "T1  -" not in block:
                continue
            def grab(tag: str) -> list[str]:
                return re.findall(rf"^{tag}\s+-\s+(.*)$", block, re.M)
            records.append({
                "year": (grab("PY") or [""])[0][:4],
                "title": " ".join(grab("T1")),
                "doi": (grab("DO") or [""])[0],
            })
    ddos = [r for r in records if re.search(r"ddos|denial[- ]of[- ]service", r["title"], re.I)]
    summary = {
        "n": len(records),
        "years": dict(sorted(Counter(r["year"] for r in records).items())),
        "ddos_titles": len(ddos),
    }
    lines = [
        "# Supplementary material S1. Literature screen",
        "",
        "Screened on 15 October 2026.",
        "Source: Journal of Information Security and Applications bibliographic records exported from ScienceDirect, publication years 2020–2026.",
        "Fields screened: title, abstract, and keywords.",
        f"Records: {summary['n']}.",
        "",
        "Publication-year counts:",
        "",
    ]
    for year, count in summary["years"].items():
        lines.append(f"- {year}: {count}")
    lines += [
        "",
        "Records with publication year 2024 are included in this count. The export does not contain publication year 2027.",
        f"Titles matching DDoS or denial-of-service, including one corrigendum: {summary['ddos_titles']}.",
        "",
        "Table 1 keeps one study for each nearby problem: threshold choice, benchmark realism, tree-based NIDPS rules, hybrid robustness, imbalance remedies, drift, and edge or federated defence. Full texts of all records were not read. The screen is a positioning aid for this empirical paper. It is not a systematic review, and PRISMA is not used.",
        "",
        "## DDoS and denial-of-service titles",
        "",
    ]
    for row in sorted(ddos, key=lambda r: (r["year"], r["title"])):
        lines.append(f"- {row['year']}. {row['title']} {row['doi']}")
    log = ROOT / "05_Supplementary" / "S1_Literature_screen.md"
    log.parent.mkdir(parents=True, exist_ok=True)
    log.write_text("\n".join(lines) + "\n", encoding="utf-8")
    return summary


def build_manuscript(summary: dict):
    doc = Document()
    for section in doc.sections:
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)

    add_para(doc, TITLE, size=14, bold=True, align="center", space_after=12)
    add_h(doc, "Abstract")
    add_para(doc, ABSTRACT, align="justify")
    add_para(doc, "Keywords: " + "; ".join(KEYWORDS), italic=True)

    add_h(doc, "1. Introduction")
    add_para(
        doc,
        "A detector that looks excellent on a public DDoS benchmark can still be expensive to operate. False alarms "
        "become alerts that an analyst or an automated control must handle, and missed attacks remain on the path "
        "to the service (Sommer and Paxson, 2010; Zargar et al., 2013). On an attack-heavy benchmark both kinds of "
        "error can hide inside an F1 score above 99%. The base-rate problem is old (Axelsson, 2000), but it is "
        "easy to lose once leaderboards report a single threshold on a single mix of traffic.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        f"The scale is the count, not a rounded percentage. On the CICIoT2023 slice that was actually scored, the "
        f"CNN–LSTM hybrid produces {NUMS['hybrid_fp']} false alarm in {NUMS['hybrid_benign']:,} benign windows, and "
        f"random forest produces {NUMS['rf_fp']} false alarms in {NUMS['rf_benign']:,} benign records. Those two "
        "denominators are not the same test. Both models still exceed 99% F1, the forest misses fewer attacks "
        f"({NUMS['rf_fnr_pct']:.2f}% versus {NUMS['hybrid_fnr_pct']:.2f}%), and the forest has the higher ROC area. "
        "Which detector is preferable therefore depends on how often attacks occur and on how the two errors are "
        "valued (Elkan, 2001; Provost and Fawcett, 2001). The counts come from the executed notebooks behind "
        "Iqbal (2026). This article does not retrain the models, and it does not choose a new threshold.",
        align="justify", first_line=True,
    )
    add_h(doc, "1.1 Questions")
    add_para(
        doc,
        "Four questions organise the paper. First, at the two-class argmax that was used, how do false-alarm rate "
        "(FAR) and miss rate (FNR) differ on the CICIoT2023 slices that were actually scored? Second, on "
        "CICDDoS2019, which is not an IoT trace, do the recovered false-alarm counts put random forest ahead of "
        "the hybrid, or only a rounded F1? Third, if the measured true-positive rate and FAR stay "
        "constant, how does precision move as attack prevalence falls? Fourth, for the hybrid and random forest, "
        "at what prevalence does the lower-cost detector change, and where do never alerting and always alerting "
        "become cheaper?",
        align="justify", first_line=True,
    )
    add_h(doc, "1.2 Contribution")
    add_para(
        doc,
        "The detectors are established, and the paper does not introduce a new layer or a new optimiser "
        "(Hochreiter and Schmidhuber, 1997; Akiba et al., 2019; Grinsztajn et al., 2022). The contribution is a "
        "security reading of operating points recovered from the executed notebooks, including the exact confusion "
        "counts and the fact that the CICIoT2023 deep and tree results are not paired. On CICDDoS2019 the "
        "recovered rates reverse that false-alarm order at the evaluated points. Equation (1) projects precision, and "
        "equation (3) is the iso-performance condition of Provost and Fawcett (2001) for these two fixed points, "
        "not a new decision algorithm. Of the eight learned CICIoT2023 points, the ROC convex hull contains only the hybrid "
        "and random forest (Fawcett, 2006). The same arithmetic marks where never alerting or always alerting is "
        "cheaper. That boundary describes the evaluated cuts. It does not describe thresholds selected to a shared "
        "false-alarm budget, and it does not describe a security-operations workflow.",
        align="justify", first_line=True,
    )

    add_h(doc, "2. Related work")
    add_h(doc, "2.1 DDoS detectors on public benchmarks")
    add_para(
        doc,
        "CICDDoS2019 and CICIoT2023 are widely used, attack-heavy benchmarks built from engineered flow and IoT "
        "traffic features (Sharafaldin et al., 2019; Neto et al., 2023; Ring et al., 2019). Deep detectors on this "
        "family include lightweight convolutional models (Doriguzzi-Corin et al., 2020), two-stage DDoS frameworks "
        "(Mittal et al., 2023), and hybrid convolutional–recurrent networks evaluated for adversarial robustness "
        "(González-Ramos et al., 2026). Sequential detectors for low-latency services make the same architectural "
        "bet, that order inside a short window carries attack evidence (Cogranne et al., 2026). Tree ensembles "
        "remain competitive when the features are already tabular, and they can be compiled into NIDPS rules "
        "(Coscia et al., 2024; Grinsztajn et al., 2022; Shwartz-Ziv and Armon, 2022). Edge, SDN, and multi-tenant "
        "deployments change the cost of a false alarm more than they change the model catalogue (Shi et al., 2016; "
        "Meidan et al., 2018; Li et al., 2025; Madoune et al., 2025; Singha et al., 2025; Anurupam and Seshadri, "
        "2025; Hormozi et al., 2026; Javadpour et al., 2025). Surveys of the area now cover both classical and "
        "neurosymbolic designs (Ferrag et al., 2020; Bizzarri et al., 2025).",
        align="justify", first_line=True,
    )
    add_h(doc, "2.2 Operating points, base rates, and cost")
    add_para(
        doc,
        "A single cut at 0.5 is only one point on a score. ROC area summarises all cuts and can rank a model "
        "differently from the cut that an operator actually uses (Saito and Rehmsmeier, 2015; Fawcett, 2006). "
        "Provost and Fawcett (2001) showed that, once class proportions and error costs are specified, the preferred "
        "operating point lies on the ROC convex hull along an iso-performance line whose slope is set by those two "
        "quantities. Elkan (2001) gives the corresponding cost-sensitive decision rule. In intrusion detection the same "
        "arithmetic appears as the base-rate fallacy: a small false-alarm rate produces mostly false alerts when "
        "attacks are rare (Axelsson, 2000). Gu et al. (2006) later argued for measures of detection capability "
        "that keep the base rate in view, rather than accuracy alone. Threshold moving is already used inside "
        "DDoS detectors (Ko et al., 2020), and some recent anomaly detectors avoid a fixed cut entirely "
        "(Zhao et al., 2026). Calibration matters because a nominal probability of 0.5 need not be the cost-optimal "
        "cut (Guo et al., 2017).",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Imbalance is often treated by changing the training distribution, for example with class-specific "
        "generative models or graph models that up-weight rare attacks (Wang et al., 2026a; Wang et al., 2026b; "
        "Lara-Gutierrez et al., 2026). Drift-aware systems ask a related question after deployment, when the mix "
        "itself moves (Rustam and Jurcut, 2025). Benchmark realism is a separate concern: synthetic attack-heavy "
        "collections need not resemble an operational mix (Layeghy et al., 2024). Evaluation pitfalls in security "
        "machine learning, including leakage and unrealistic base rates, are set out by Sommer and Paxson (2010), "
        "Arp et al. (2022), and Apruzzese et al. (2023). Alert volume after the detector has fired is itself an "
        "operational task (Vaarandi and Guerra-Manzanares, 2024).",
        align="justify", first_line=True,
    )
    add_h(doc, "2.3 Positioning")
    add_para(
        doc,
        f"Supplementary material S1 describes a screen of {summary['n']} Journal of Information Security and "
        "Applications records from 2020 to 2026. Table 1 keeps one study for each neighbouring problem. The screen "
        "is a positioning aid. It is not a claim that the same crossover has never been published, and it is not a "
        "systematic review. PRISMA does not apply.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The neighbouring papers typically contribute a detector, a defence, or a rebalancing method. This paper "
        "does not. It also does not do the next operational experiments those papers leave open: it does not select "
        "thresholds on validation data to a shared false-alarm or alert budget (Ko et al., 2020; Provost and "
        "Fawcett, 2001), and it does not test the resulting policy in a security-operations workflow (Vaarandi and "
        "Guerra-Manzanares, 2024). Equation (3) states where the executed hybrid and random-forest points exchange "
        "expected cost if their rates stay fixed. That is a narrower claim than a matched-threshold comparison.",
        align="justify", first_line=True,
    )

    add_h(doc, "3. Threat model and protocol")
    add_h(doc, "3.1 What the detector is asked to do")
    add_para(
        doc,
        "The defender observes engineered features and marks each record, or each window, as attack or benign. A "
        "false alarm can become a blocked flow or an alert. A miss leaves attack traffic unmarked. No control loop "
        "is simulated. Evasion, in which an attacker perturbs features to cross the cut, is outside the experiments "
        "(ElShehaby and Matrawy, 2026; González-Ramos et al., 2026). Edge hardware is part of the deployment story "
        "(Shi et al., 2016; Anurupam and Seshadri, 2025) and is not part of the measurements: latency and memory "
        "were not taken on a device.",
        align="justify", first_line=True,
    )
    add_h(doc, "3.2 Data")
    add_para(
        doc,
        "Table 2 lists the counts printed by the executed notebooks behind Iqbal (2026). CICDDoS2019 began from the "
        "official daily CSV files (CIC, 2019). Within each file, attack rows were capped at five times the benign "
        "rows, with the sample drawn at seed 42, and the files were concatenated. The loaded table has 630,457 "
        "rows. After a timestamp sort, a random stratified split (seed 42) assigned 70% to training, 20% to "
        "validation, and 10% to test: 441,319, 126,092, and 63,046 records, with 82 features. Each split is 16.91% "
        "benign. CICIoT2023 was loaded from a dissertation archive that already contained separate training, "
        "validation, and test files, rather than from the merged CSVs distributed by CIC. Those files have "
        "5,491,971, 1,176,851, and 1,176,851 records and 46 features. The test file contains 27,709 benign and "
        "1,149,142 attack records. The training file contains 129,538 benign and 5,362,433 attack records, so "
        "the training attack share is 97.64% (CIC, 2023; Sharafaldin et al., 2019; Neto et al., 2023).",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Classical models are scored on records. Deep models are scored on windows of 100 records, stride 1, whose "
        "label is the label of the last record. Windows are built inside each split. On CICDDoS2019 every test "
        "window was scored: 62,947 windows, of which 10,650 are benign, against 63,046 records from the same "
        "split. On CICIoT2023 the test file yields 1,176,752 windows, but the notebook scores only the first "
        "500 batches of 512, which is 256,000 windows. Those windows are 2.33% benign, close to the 2.35% benign "
        "share of the whole test file. The file has no timestamp, address, or flow identifier, so a window of "
        "100 consecutive rows is not a verified time series. The classical CICIoT2023 models are scored on a "
        "different draw: a stratified sample of 500,000 training records and 300,000 test records, seed 42.",
        align="justify", first_line=True,
    )
    add_h(doc, "3.3 Preprocessing")
    add_para(
        doc,
        "Figure 1 shows the two pipelines that produced the tables, which are not the same pipeline. On CICDDoS2019 "
        "the raw timestamp, source and destination IP, Flow ID, and SimillarHTTP fields were removed, and a numeric "
        "Unix timestamp was added back as a feature. Missing values were then replaced by medians computed on all "
        "630,457 rows, before the split, so the test rows influenced the imputer. The rows were sorted by time and "
        "then split at random, so the sort does not give a chronological test. On CICIoT2023 every column except "
        "the label was kept, infinite values were treated as missing, and missing values were filled with zero "
        "inside each split. In both notebooks a min–max scaler was fit on training features only and applied to "
        "validation and test (Kapoor and Narayanan, 2023). The validation split selected the hybrid’s "
        "hyperparameters. Section 5 uses the test scores from the slices named above.",
        align="justify", first_line=True,
    )
    add_h(doc, "3.4 Metrics and the decision boundary")
    add_para(
        doc,
        "FAR is the share of benign cases marked as attack, and FNR is the share of attack cases marked as benign. "
        "Recall, the true-positive rate, is 1 − FNR. Accuracy, precision, F1, ROC area, and precision–recall area "
        "are reported where the source tables include them. On CICDDoS2019 the Matthews correlation is the more informative "
        "summary, because several F1 values round to 100.00 (Chicco and Jurman, 2020).",
        align="justify", first_line=True,
    )
    add_rich(
        doc,
        [
            ("If a detector keeps the same true-positive rate and FAR when the attack prevalence changes to a deployment value ", False),
            ("π", True),
            (", precision is", False),
        ],
        first_line=True, space_after=2,
    )
    add_rich(
        doc,
        [("(1)    Precision(", False), ("π", True), (") = (TPR × ", False), ("π", True), (") / (TPR × ", False), ("π", True), (" + FAR × (1 − ", False), ("π", True), (")).", False)],
        align="center", space_after=4,
    )
    add_para(
        doc,
        "Equation (1) is Axelsson’s base-rate identity applied to one operating point, with TPR and FAR held constant: "
        "a projection, not a resample from a rarer attack stream. The expected cost per case at the same point, with "
        "false-alarm cost cFP and miss cost cFN, is",
        align="justify", first_line=True, space_after=2,
    )
    add_rich(
        doc,
        [("(2)    ", False), ("C", True), (" = ", False), ("c", True), ("FP × FAR × (1 − ", False), ("π", True), (") + ", False), ("c", True), ("FN × FNR × ", False), ("π", True), (".", False)],
        align="center", space_after=4,
    )
    add_para(
        doc,
        "Write r = cFN / cFP. For the hybrid (H) and random forest (RF) at the evaluated cut, FARH is smaller than "
        "FARRF and FNRH is larger than FNRRF. Let ΔFAR = FARRF − FARH and ΔFNR = FNRH − FNRRF, using the exact "
        "evaluated fractions rather than percentages rounded to two decimal places. Substituting the two points into "
        "equation (2) and comparing the costs shows that the hybrid has the lower expected cost of the two exactly "
        "when the attack prevalence lies below",
        align="justify", first_line=True, space_after=2,
    )
    add_rich(
        doc,
        [("(3)    ", False), ("π", True), ("*(", False), ("r", True), (") = ΔFAR / (ΔFAR + ", False), ("r", True), (" × ΔFNR).", False)],
        align="center", space_after=4,
    )
    add_para(
        doc,
        "Above that prevalence, random forest has the lower expected cost of the two learned detectors. Equation (3) "
        "is the iso-performance condition of Provost and Fawcett (2001) specialised to two fixed points rather than "
        "to a whole ROC curve. It is not a new algorithm. Two trivial policies bound the corners. Never alerting, "
        "which has FAR = 0 and FNR = 1, costs less than the hybrid when π is below FARH / (FARH + r × TPRH). Always "
        "alerting, which has FAR = 1 and FNR = 0, costs less than random forest when π is above "
        "(1 − FARRF) / (1 − FARRF + r × FNRRF). For every cost ratio used below, these breakpoints are ordered, so "
        "the cheapest policy among never alerting, the hybrid, random forest, and always alerting changes in that "
        "sequence as attack prevalence rises. That order belongs to these two unequal slices, not to full ROC "
        "curves and not to a shared test set. The cost ratios used below (0.1, 1, 10, and 100) are scenarios, for "
        "example a short analyst review against a much costlier outage. They are not measured prices from a "
        "security operations centre. An alert budget gives the same quantities an operational reading: if at most "
        "A false alerts can be handled among B benign flows, the allowable FAR is A/B, and the relevant comparison "
        "is the miss rate at that FAR. That comparison is not in Section 5.",
        align="justify", first_line=True,
    )

    add_h(doc, "4. Models")
    add_para(
        doc,
        "The classical models are logistic regression, Gaussian naive Bayes, a decision tree, random forest, and "
        "either extra trees on CICDDoS2019 or histogram gradient boosting on CICIoT2023, implemented in "
        "scikit-learn (Pedregosa et al., 2011). The deep models are a one-dimensional convolutional network, a "
        "long short-term memory network, and a hybrid that stacks the convolution with a recurrent layer on a "
        "window of length 100 and emits two logits (Hochreiter and Schmidhuber, 1997). The reported decision is "
        "the argmax of those logits, which is the same cut as a softmax probability of one half except on exact "
        "ties. It is not a sigmoid followed by binary cross-entropy. The loss is class-weighted cross-entropy. "
        "On CICIoT2023 the benign and attack weights are the balanced values 21.20 and 0.51. Logistic regression "
        "and the decision tree use class_weight balanced, and random forest uses balanced subsampling. Naive Bayes "
        "and histogram gradient boosting are unweighted. The same pattern, with extra trees weighted and naive "
        "Bayes unweighted, is used on CICDDoS2019. The FAR gap is therefore not a contrast between a weighted "
        "network and an unweighted forest.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The hybrid’s width was selected by the tree-structured Parzen estimator (Bergstra et al., 2011; Snoek et al., "
        "2012), implemented in Optuna (Akiba et al., 2019). On CICIoT2023 the search was 10 trials of 3 epochs, "
        "capped at 300 training batches and 150 validation batches per epoch, and the objective was validation "
        "accuracy. That objective is weak when 97.64% of the training rows are attacks. The selected setting is 128 convolutional filters, a recurrent width of 100, 3 layers, learning "
        "rate 0.00205, and dropout 0.179. The final hybrid was then trained for 15 epochs, but each epoch used only "
        "1,200 of the 10,727 training batches. The convolutional and recurrent baselines used the same batch cap "
        "for 8 epochs, with 64 filters or a width of 100 and 2 layers. On CICDDoS2019 the search was 20 trials of "
        "5 epochs, capped at 200 training batches, maximising validation F1. The selected setting is 64 filters, "
        "width 50, 1 layer, learning rate 0.00110, dropout 0.319, and weight decay 3.6×10−6; the final deep models "
        "on that dataset were not batch-capped. Classical models were not given this search. The CICIoT2023 forest "
        "was fit on 500,000 stratified training rows. The CICDDoS2019 trees were fit on 300,000 stratified training "
        "rows and scored on the full test split.",
        align="justify", first_line=True,
    )

    add_h(doc, "5. Results")
    add_h(doc, "5.1 CICDDoS2019")
    add_para(
        doc,
        "Table 3 reports false-alarm and miss counts recovered from the printed precision, recall, Matthews "
        "correlation, and class support. For every model there is one pair of non-negative integers that "
        "reproduces those six-decimal figures. Random forest has 0 false alarms in 10,663 benign records and 4 "
        "misses in 52,383 attack records. The hybrid has 19 false alarms in 10,650 benign windows and 15 misses "
        "in 52,297 attack windows. The forest therefore has both the lower false-alarm rate and the lower miss "
        "rate, so equation (3) has no crossover on these two points: the forest has the lower expected cost at "
        "every prevalence and every positive cost ratio. Figure 2 plots the two false-alarm rates with 95% "
        "binomial intervals. Those intervals do not overlap when each window is treated as an independent trial. "
        "Overlapping windows and repeated flows make that assumption too strong. The units still differ, records against windows from "
        "the same split, but the order is the reverse of the CICIoT2023 order in Section 5.2. F1 values that "
        "round to 100.00 hide this. The Matthews correlations remain 0.9998 for the forest and 0.9981 for the "
        "hybrid.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The same notebook prints the fitted feature importances. The decision tree assigns 0.679 of its "
        "importance to the Inbound flag, 0.297 to the source port, and 0.014 to the destination port. Inbound "
        "also has an absolute correlation of 0.867 with the binary label. Random forest assigns 0.112 to Inbound, "
        "after 0.123 on the destination port, and 0.016 to the retained Unix timestamp. Extra trees assigns 0.308 "
        "to Inbound and 0.007 to that timestamp. Ports, a direction flag, and a timestamp can encode the "
        "capture rather than the attack (Engelen et al., 2021; Flood et al., 2024; Arp et al., 2022). These "
        "figures are importances from the fitted trees on that matrix. Section 5.6 refits a random forest after "
        "those fields are removed, on a four-file subset rather than on this matrix. The medians here were still "
        "fit before the split, so the near-ceiling scores are not a cleaned result. "
        "CICDDoS2019 is not an IoT trace. Table 4 lists the recorded fit times. Classical fits take from a "
        "fraction of a second to about 21 seconds. Deep-model fit times for this dataset were not printed. The "
        "times are notebook wall-clock times, not latencies on a gateway.",
        align="justify", first_line=True,
    )
    add_h(doc, "5.2 CICIoT2023 at the evaluated cut")
    add_para(
        doc,
        "Table 5 rounds the metrics to two decimal places. Table 6 is the confusion matrix the notebook printed, "
        "and equations (1)–(3) use those integers. The hybrid’s false-alarm count is 1 of 5,962 benign windows, "
        f"which is {NUMS['hybrid_far_pct']:.3f}% and rounds to the 0.02% in Table 5. Its miss count is "
        f"{NUMS['hybrid_fn']:,} of 250,038 attack windows ({NUMS['hybrid_fnr_pct']:.2f}%). Random forest’s "
        f"false-alarm count is {NUMS['rf_fp']} of {NUMS['rf_benign']:,} benign records "
        f"({NUMS['rf_far_pct']:.2f}%) and its miss count is {NUMS['rf_fn']} of 292,936 "
        f"({NUMS['rf_fnr_pct']:.2f}%). The ratio of the two exact false-alarm rates is "
        f"{NUMS['far_fold']:.0f}-fold. The older summary of roughly 300-fold came from rounding 0.017% up to 0.02% "
        "before dividing. A binomial 95% interval for the hybrid’s one false positive, treating the 5,962 windows "
        f"as independent, runs from {NUMS['far_ci_low_pct']:.4f}% to {NUMS['far_ci_high_pct']:.3f}%. Overlapping "
        "windows of stride 1 violate that independence, so the interval is too narrow rather than too wide "
        "(Pendlebury et al., 2019). The point ratio is not a stable fold.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The hybrid’s ROC area is 99.63% and the forest’s is 99.96%, so the model with fewer false alarms at this "
        "cut is not the model with the higher ROC area. Histogram gradient boosting is close to the forest on the "
        "same 300,000 records (427 false alarms and 779 misses). The convolutional network has 5 false alarms in "
        "the same 5,962 windows as the hybrid’s one false alarm, and the binomial intervals overlap, so this cut "
        "does not separate the hybrid from that baseline. Figure 3 plots the eight pairs. The black curve "
        "connects the hybrid and random forest, the learned vertices of the convex hull of these fixed points plus "
        "never alerting and always alerting. It is not a threshold sweep. LSTM lies just inside that edge. The bar "
        "on the hybrid is the binomial interval above, not a repeated-run interval. The dashed line is the "
        "false-alarm rate at which, under 1% attack prevalence and equal assumed error costs, a detector matches "
        "the hybrid’s expected cost even if it misses nothing. The tree points lie well to the right of that line. "
        "They also lie on a different sample, so the horizontal distance is not a paired difference.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Table 6 is not a rescaling. It is the confusion matrix printed for each model on the slice that model was "
        "scored on: 256,000 windows for the three deep models, and 300,000 records for the five classical models. "
        "The deep models share the 5,962 benign windows. The classical models share the 7,064 benign records. "
        "Applying the 2.36% training benign share to all 1,176,851 test records, which is what a rounded "
        "illustration would do, does not recover these cells. The hybrid’s single false positive would be "
        "overstated by that illustration, and the forest’s 431 false positives would be assigned to the wrong "
        "denominator. No significance test is reported. One seed was used, and overlapping windows are not "
        "independent (Pendlebury et al., 2019).",
        align="justify", first_line=True,
    )
    add_h(doc, "5.3 Precision when prevalence changes")
    add_para(
        doc,
        "Table 7 and Figure 4 apply equation (1) on a dense prevalence grid, using the exact rates from Table 6. "
        "At the training attack share of 97.64%, projected precision stays near the benchmark precision. "
        f"At 1% attack prevalence the projection is {NUMS['hybrid_precision_1pct']:.2f}% for the hybrid and "
        f"{NUMS['rf_precision_1pct']:.2f}% for random forest. At 0.1% it is "
        f"{NUMS['hybrid_precision_0_1pct']:.2f}% and {NUMS['rf_precision_0_1pct']:.2f}%. Random forest and "
        "histogram gradient boosting remain distinct lines because their false-alarm rates differ slightly. The "
        "curves hold TPR and FAR constant. Changing the attack prevalence while holding those conditional rates "
        "fixed is an analytical step. It is not evidence that the same rates would be observed on a network whose "
        "devices, attackers, or timing also differ (Layeghy et al., 2024; Rustam and Jurcut, 2025).",
        align="justify", first_line=True,
    )
    add_h(doc, "5.4 Where the lower-cost detector changes")
    add_para(
        doc,
        f"For the exact rates, ΔFAR = {NUMS['delta_far_pct']:.2f} percentage points and ΔFNR = "
        f"{NUMS['delta_fnr_pct']:.2f} percentage points. Equation (3) then gives one crossover prevalence for each "
        f"cost ratio. At equal assumed costs, r = 1, the crossover is {NUMS['crossover_equal_pct']:.2f}%. "
        "Of the two learned detectors, random forest has the lower expected cost when attacks are more common "
        "than the crossover, which includes the CICIoT2023 training mix of 97.64%. The hybrid has the lower "
        "expected cost of the two when attacks are rarer, which includes 1% prevalence. Table 8 lists the "
        "boundary at r = 0.1, 1, 10, and 100, together with the trivial-policy bounds. Figure 5 draws the "
        "hybrid–forest boundary as a continuous curve and shades the regions in which never alerting or always "
        "alerting is cheaper. At r = 100, always alerting is cheaper than random forest once attack prevalence "
        f"exceeds {100 * dm.always_alert_prevalence(100.0, RF['FAR'] / 100, RF['FNR'] / 100):.1f}%, so the "
        "training mix does not make the forest the cheapest policy when misses are priced that far above false "
        "alarms. These boundaries are conditional on the evaluated slices and on the assumed cost ratios.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The same rates convert into an alert load only under a stated volume and only if the false-alarm rate "
        f"travels. At ten million benign flows, the hybrid’s exact rate is {NUMS['alerts_hybrid']:,.0f} false "
        f"alerts and random forest’s is {NUMS['alerts_rf']:,.0f}. The hybrid figure is one false positive scaled "
        f"up. The binomial interval on that count, still assuming independence, spans about "
        f"{NUMS['alerts_hybrid_ci_low']:,.0f} to {NUMS['alerts_hybrid_ci_high']:,.0f} alerts per ten million benign "
        f"flows, and dependence would widen it (Vaarandi and Guerra-Manzanares, 2024). If the attack prevalence "
        f"is 1%, that benign volume corresponds to about {ATTACKS_AT_1PCT:,.0f} attacks, of which the hybrid "
        f"misses about {MISS_H:,.0f} and random forest misses about {MISS_RF:,.0f}. At the evaluated cuts and "
        "with equal unit costs, the hybrid is the cheaper of the two learned points, because the forest’s much "
        "larger false-alarm rate dominates the cost. A deployment that prices a miss far above a false alarm "
        "moves back toward the forest, until the always-alert region in Figure 5 is reached. None of these counts "
        "was observed in a security operations centre.",
        align="justify", first_line=True,
    )

    add_h(doc, "5.5 BoT-IoT on a complete mirror subsample")
    add_para(
        doc,
        "Table 9 and Table 10 report a within-dataset run on a Kaggle mirror of the UNSW BoT-IoT 5% all-features "
        "extract (Koroniotis et al., 2019; Kaggle dataset vigneshvenkateswaran/bot-iot-5-data). It is not a transfer "
        "of the CIC models, and the CNN–LSTM was not retrained. The mirror parses completely. Its byte lengths differ "
        "from the published CloudStor lengths, so it is a complete parseable mirror rather than a byte-verified copy "
        "of the official files. The third file is 285,316,035 bytes, against a published length of 282,131,555. "
        "Every one of the 477 normal rows is kept. Attack rows are subsampled so that the modelled table has "
        "350,000 rows: 183,869 DDoS, 157,034 DoS, 8,610 reconnaissance, 477 normal, and 10 theft. Those category "
        "counts are the subsample, not the published 5% census.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Addresses, timestamps, the row index, and the sequence field were removed before the split. Medians and the "
        "min–max scaler were fit on training rows only. The split is stratified, with 70% for training and 15% each "
        "for validation and test. Every model is scored on those same records. The audited features are the flow "
        "counters and the numeric protocol, state, and flag codes. Two further sets add the source and destination "
        "ports, or the published IP-grouped aggregate fields, but not the raw addresses. The models, fit in "
        "scikit-learn (Pedregosa et al., 2011), are a depth-capped decision tree, a random forest of 80 trees "
        "drawing bags of 150,000 rows, histogram gradient boosting, and class-weighted logistic regression. The "
        "forest is not fit on every training row. Once the training set exceeds 180,000 rows, logistic regression "
        "is fit on at most 40,000 benign rows and 140,000 attack rows. Seeds 42, 7, and 123 are used for the audited "
        "flow features. The other feature sets use seed 42, which redraws the same split.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Seed 42 has 72 benign test records and 71 benign validation records, so one false alarm is about 1.4% of "
        "the benign test rows. A 0.1% budget and a 0.02% budget both permit zero false positives on 71 benign "
        "validation records, and they therefore cannot support a precise sub-0.1% false-alarm claim. Table 9 uses "
        "the 0.5 cut. The decision tree has 0 false alarms and 175 misses at seed 42, 0 and 201 at seed 7, and 1 "
        "and 188 at seed 123. Moving the seed-42 tree’s threshold to the validation rule that allows one false "
        "alarm leaves the test result at 0 false alarms and 175 misses. Seeds 7 and 123 do change: the tree moves "
        "from 0 false alarms and 201 misses to 2 and 75, and from 1 false alarm and 188 misses to 1 and 73. The "
        "same seed-42 rule lowers the forest from 0 false alarms and 168 misses to 1 and 49, and histogram boosting "
        "from 1 and 48 to 1 and 47. Logistic regression’s tied scores leave two validation false alarms, and its "
        "test misses rise from 288 to 34,380.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "Table 10 holds the split at seed 42. At the 0.5 cut, restoring ports leaves the tree at 0 false alarms and "
        "135 misses, against 175 on flow features alone, and histogram boosting falls from 48 misses to 8, with "
        "false alarms rising from 1 to 2. On this split the validation cut does not reduce the tree’s misses, while "
        "restoring ports does. The forest’s tuned cut does reduce its misses. When the IP-grouped aggregates are "
        "added at the 0.5 cut, the decision tree assigns 0.965 of its importance to N_IN_Conn_P_DstIP and has 0 "
        "false alarms and 38 misses. The forest and histogram boosting then have 1 false alarm and 5 misses each. "
        "Those fields are published per-address aggregates; this run does not establish that they are available at "
        "prediction time. Figure 6 plots the seed-42 counts. Threshold, features, and seed still interact, and the "
        "interaction is not the one measured on the earlier truncated download. These controls do not replace the "
        "CICIoT2023 or CICDDoS2019 comparisons in Sections 5.1 and 5.2.",
        align="justify", first_line=True,
    )
    add_h(doc, "5.6 A four-file CICDDoS2019 shortcut check")
    add_para(
        doc,
        "Table 11 is a separate random-forest check on 120,444 rows rebuilt from four CICDDoS2019 day files "
        "(UDPLag, Syn, NTP, and LDAP), using the same five-to-one attack cap as the original loading script and "
        "one seed. It is not the 630,457-row matrix behind Table 3, and the hybrid was not refit. The timestamp "
        "column in this rebuild is text and is already excluded from the numeric features, so dropping it changes "
        "nothing. Exact duplicate rows number 0. Flow identifiers overlap the training and test splits in 1,869 "
        "cases, and source–destination pairs overlap in 605. Medians for the main rows were fit on training rows "
        "only. At the 0.5 cut the forest has 1 false alarm and 3 misses in 3,011 benign and 15,056 attack test "
        "records. Removing Inbound leaves the false alarms at 1 and raises the misses to 10. Removing Inbound and "
        "both port fields raises the counts to 5 false alarms and 18 misses. That false-alarm count, 5 of 3,011, "
        "is about 0.17%. The historical hybrid rate in Table 3 is 19 of 10,650 windows, about 0.18%, on a different "
        "population. The two rates are close enough that the CICDDoS2019 side of the reversal needs a paired refit "
        "before it can be treated as robust to these fields. A group-held-out split was not run.",
        align="justify", first_line=True,
    )

    add_h(doc, "6. Discussion")
    add_h(doc, "6.1 How to read the boundary")
    add_para(
        doc,
        "Equations (1)–(3), the ROC hull of the evaluated points, and the trivial-policy bounds support the reading "
        "in Box 1. The calculations in this article carry out the prevalence and cost steps at the cut that was "
        "used. Step 2, choosing the threshold on validation data, is the comparison an operator should run next. "
        "Random forest’s higher ROC area is a reason to expect that its miss rate at a matched false-alarm budget "
        "may differ from its miss rate at this cut, which could move or remove the crossover. The two points are "
        "also not paired, so a matched budget has to be computed on aligned cases before the ranking can be "
        "treated as a property of the architectures.",
        align="justify", first_line=True,
    )
    add_procedure(doc)
    add_para(
        doc,
        "On CICDDoS2019 the recovered counts give random forest both a lower false-alarm rate and a lower miss "
        "rate than the hybrid, so that dataset has no hybrid–forest crossover. On CICIoT2023 the evaluated cut "
        "separates them in the other direction, on cases that are not shared. Detector selection is then a "
        "statement about the dataset, the scored cases, and the prevalence and cost ratio. It is not a statement "
        "that one architecture is uniformly better.",
        align="justify", first_line=True,
    )
    add_h(doc, "6.2 Scope of the evidence")
    add_para(
        doc,
        "The comparisons in Section 5 come from one training seed, so they have no repeated-run interval. The bar in "
        "Figure 3 is a binomial interval on the printed counts. For the deep models it treats overlapping windows "
        "as independent and is therefore too narrow. The hybrid’s one false positive can move the fold; it moves "
        "the equal-cost crossover much less, because that boundary is set mainly by the forest’s 431 false "
        "positives. Deep and classical CICIoT2023 metrics use different cases: the first 256,000 test windows "
        "versus a stratified sample of 300,000 test records. On CICDDoS2019 the units differ as well, the medians "
        "were fit before the split, a Unix timestamp was retained, and the split is random after sorting, so "
        "duplicate, device, and campaign overlap were not measured and are not claimed to be absent (Arp et al., "
        "2022; Kapoor and Narayanan, 2023; Pendlebury et al., 2019). Class weights and tuning budgets are now "
        "known and are not matched across families. Each CICIoT2023 deep epoch used 1,200 of 10,727 training "
        "batches. Equations (1) and (3) hold TPR and FAR constant when prevalence changes, and the cost ratios "
        "are scenarios rather than measured operational prices. Section 5.5 is a BoT-IoT subsample on a Kaggle mirror, "
        "not a transfer from the CIC splits (D’hooge et al., 2020). There is no measurement "
        "on an edge device. Softmax scores were computed in the notebooks and used for the ROC areas, but the original "
        "CIC score files and checkpoints were not in the recovered archive. The original CICIoT2023 slices remain unpaired. "
        "Supplementary material S2 records a later paired pilot on a balanced benign-versus-ICMP-flood subset from a "
        "Hugging Face copy of the raw files. Both models miss no attacks on that subset, so the pilot demonstrates the "
        "protocol and does not test the comparison in Section 5.2, nor the other attack families in CICIoT2023. "
        "Wolf et al. (2026) likewise find that a classifier fit on emulated traffic need not match one fit on the "
        "corresponding hardware capture, so one benchmark’s operating point is not a deployment result. "
        "On BoT-IoT, the 0.1% and 0.02% budgets "
        "are the same rule, because 71 benign validation records cannot resolve them.",
        align="justify", first_line=True,
    )

    add_h(doc, "7. Conclusion")
    add_para(
        doc,
        "On CICDDoS2019, random forest has 0 false alarms in 10,663 benign records and the hybrid has 19 in "
        "10,650 benign windows, so the forest dominates on the recovered rates and there is no cost crossover. "
        "On CICIoT2023 the order reverses at the evaluated cut: the hybrid’s false-alarm count is "
        f"{NUMS['hybrid_fp']} of {NUMS['hybrid_benign']:,} windows and random forest’s is {NUMS['rf_fp']} of "
        f"{NUMS['rf_benign']:,} records. Those CICIoT cases are different. Held fixed, the CICIoT rates project "
        f"to a precision of {NUMS['hybrid_precision_1pct']:.1f}% and {NUMS['rf_precision_1pct']:.1f}% at 1% "
        f"attack prevalence, and equal assumed costs switch the cheaper detector at "
        f"{NUMS['crossover_equal_pct']:.1f}% prevalence. A decision tree on CICDDoS2019 concentrates about 99% "
        "of its importance on an Inbound flag and two ports, and the other trees also put substantial weight on "
        "Inbound. On a four-file subset, removing that flag and the ports raises a random forest from 1 false alarm "
        "and 3 misses to 5 and 18, so the CICDDoS2019 advantage is not yet shown to survive those fields. A BoT-IoT "
        "mirror subsample that keeps all 477 benign rows does not repeat an earlier drop in tree misses: at seed 42 "
        "the tuned cut stays at 175 misses, while seeds 7 and 123 do drop. It does not rebuild the "
        "CIC comparisons. Rankings of these detectors should name the "
        "dataset, the scored cases, and the prevalence and cost ratio.",
        align="justify", first_line=True,
    )

    add_h(doc, "Declaration of competing interests")
    add_para(doc, "The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper.", align="justify")
    add_h(doc, "Funding")
    add_para(doc, "This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.", align="justify")
    add_h(doc, "Data availability")
    add_para(
        doc,
        "CICDDoS2019 and CICIoT2023 are available from the Canadian Institute for Cybersecurity (CIC, 2019, 2023). "
        "The historical CIC counts in Tables 3 and 6, and the protocol in Section 3, are those printed by the "
        "executed notebooks behind Iqbal (2026). Derived tables and scripts for those calculations are deposited "
        "at https://doi.org/10.5281/zenodo.23435601 (Iqbal, Premebida and Haider, 2026). That interim record still "
        "contains an earlier BoT-IoT run on a truncated third file. Tables 9 and 10 in this manuscript replace that "
        "run. They come from a Kaggle mirror of the BoT-IoT 5% extract (Koroniotis et al., 2019), subsampled to "
        "350,000 rows while keeping all 477 benign rows. The raw files are not redistributed. Table 11 comes from "
        "a four-file CICDDoS2019 subset. Supplementary material S2 comes from a benign-versus-ICMP-flood pilot and "
        "is not a reanalysis of Table 6. Iqbal (2026) is available from the corresponding author on request. The "
        "original CIC score files, split indices, and model checkpoints are not included.",
        align="justify",
    )
    add_h(doc, "CRediT authorship contribution statement")
    add_para(doc, "H. Iqbal: Conceptualization; Methodology; Software; Formal analysis; Investigation; Data curation; Visualization; Writing – original draft; Project administration. C. Premebida: Supervision; Writing – review & editing; Validation. U. Haider: Writing – review & editing; Validation.", align="justify")
    add_h(doc, "Declaration of generative AI and AI-assisted technologies in the manuscript preparation process")
    add_para(
        doc,
        "During the preparation of this work the authors used ChatGPT and Cursor for drafting and editing text, "
        "for restructuring sections, for assistance with plotting code, and for writing the BoT-IoT, CICIoT2023 "
        "pilot, and CICDDoS2019 subset scripts. The authors reviewed those scripts and executed them. The historical "
        "CIC prevalence, crossover, and alert figures were computed from the confusion counts printed by the executed "
        "notebooks. Tables 9 and 10 come from the mirror subsample. Table 11 comes from the four-file subset. "
        "Supplementary material S2 comes from the benign-versus-ICMP-flood pilot. The tools did not produce the "
        "original CICIoT2023 paired comparison, a five-seed refit, or a hybrid refit on the reduced CICDDoS2019 "
        "features. The authors reviewed the manuscript and take full responsibility for it.",
        align="justify",
    )

    add_h(doc, "References")
    for item in refs():
        p = add_para(doc, item, size=10, space_after=2, double=False)
        p.paragraph_format.left_indent = Inches(0.25)
        p.paragraph_format.first_line_indent = Inches(-0.25)

    add_h(doc, "Figure captions")
    for caption in [
        "Figure 1. The two executed pipelines. CICDDoS2019 imputes medians before a random split and keeps a Unix timestamp. CICIoT2023 fills missing values with zero inside each supplied split. Both fit a min–max scaler on training rows only. Deep models use a window of 100 rows labelled by the last row. The CICIoT2023 deep metrics use the first 256,000 test windows; the tree metrics use a separate sample of 300,000 test records.",
        "Figure 2. False-alarm rates at the evaluated points, with 95% Clopper–Pearson intervals. On CICDDoS2019 the models share a split, but trees are scored on records and the hybrid on windows. The forest’s observed rate is zero, so the marker is the upper end of its interval. On CICIoT2023 the two rates come from different cases and one cut. The right-hand axis is logarithmic, and the forest interval is labelled beside the marker because it is narrow on that axis. Overlapping windows make the deep-model intervals too narrow.",
        "Figure 3. CICIoT2023 operating points at the reported argmax cut. This is one cut on unequal slices, not a threshold sweep. The black curve connects the CNN–LSTM hybrid and random forest, the learned vertices of the convex hull of these fixed points plus never alerting and always alerting. LSTM lies just inside that edge. Grey points lie inside the hull. The horizontal bar is a 95% binomial interval on the false-alarm count. For the deep models, overlapping windows make that interval too narrow. The dashed line is the false-alarm rate at which expected cost equals the hybrid’s cost when attack prevalence is 1%, assumed error costs are equal, and the miss rate is zero.",
        "Figure 4. Projected precision from equation (1). True-positive rate and false-alarm rate are held at the exact Table 6 values. The dotted line marks 1% prevalence and the dashed line marks the CICIoT2023 training prevalence of 97.64%. The prevalence axis stops at 100%. The curves are projections, not measurements at the new prevalence.",
        "Figure 5. Expected-cost regions under assumed cost ratios, with true-positive rate and false-alarm rate held at the Table 6 counts. The black curve is equation (3) for these two points, not a claim about every threshold and not a measured operational cost. Below it the hybrid has lower expected cost than random forest; above it, random forest has the lower cost of those two learned detectors. The never-alert region is the grey band and is visible in the inset. The upper band is where always alerting is cheaper than random forest. Marker labels are the crossover values in Table 8. The dotted line is the CICIoT2023 training attack prevalence (97.64%). Cost ratios are assumptions.",
        "Figure 6. BoT-IoT seed 42 on the 350,000-row mirror subsample. The left panel is missed attacks, on a logarithmic scale. The right panel is false alarms among 72 benign test records; a missing bar is a count of zero. “Flow, tuned” is the flow-only model at the validation rule that allows one false alarm. On this split that rule leaves the decision tree at 175 misses. “Ports, 0.5” and “IP fields, 0.5” restore those fields on the same split and keep the 0.5 cut. Addresses and timestamps stay out.",
    ]:
        add_para(doc, caption, size=10, space_after=3, double=False)

    add_h(doc, "Tables")
    embed_table(
        doc, load_csv("Table1_JISA_positioning.csv"),
        "Table 1. Studies used to position the operating-point reading. The screen is described in Supplementary material S1.",
        "One study is kept for each nearby problem. This is not a systematic review.",
    )
    embed_table(
        doc, load_csv("Table2_Datasets.csv"),
        "Table 2. Dataset counts, the slice that Section 5 scores, and the feature set.",
        "CICDDoS2019 attack rows were capped at five benign-row multiples inside each official daily file. CICIoT2023 was loaded from a pre-split archive, not from CIC’s merged CSVs. Deep and tree metrics on CICIoT2023 use different samples.",
        ["Dataset", "Train_records", "Validation_records", "Test_records", "Scored_in_Section_5", "Features", "Class_mix", "Source"],
    )
    embed_table(
        doc, load_csv("Table3_CICDDoS2019_metrics.csv"),
        "Table 3. False-alarm and miss counts on CICDDoS2019, recovered from the printed precision, recall, Matthews correlation, and class support.",
        "Each row is the unique non-negative integer solution that reproduces the printed six-decimal precision, recall, and Matthews correlation. The FAR interval is a 95% Clopper–Pearson interval and treats the scored rows or windows as independent. Overlapping windows make the deep-model intervals too narrow. Deep rows are windows and tree rows are records from the same split. These counts are not a saved confusion-matrix file. Table 11 is a later four-file subset, not a refit of this matrix.",
    )
    embed_table(
        doc, load_csv("Table4_Training_times.csv"),
        "Table 4. Training time, in seconds, for the runs whose durations were recorded.",
        "These are notebook wall-clock times on the machine that trained the models. Inference times cover the scored slice only. They are not latencies on an edge device. Deep-model fit times on CICDDoS2019 were not printed.",
    )
    embed_table(
        doc, load_csv("Table5_CICIoT2023_metrics.csv"),
        "Table 5. CICIoT2023 metrics at the reported argmax cut, rounded to two decimal places in percent.",
        "FAR and FNR round the fractions in Table 6. A precision of 100.00 is that rounding. Deep and classical rows use different cases.",
    )
    embed_table(
        doc, load_csv("Table6_Evaluated_counts.csv"),
        "Table 6. Confusion counts printed for the CICIoT2023 slices that were scored.",
        "Deep rows are the first 256,000 test windows. Classical rows are a stratified sample of 300,000 test records. FAR and FNR are the integer ratios, shown to four decimal places in percent. Equations (1)–(3) use the integers.",
    )
    embed_table(
        doc, load_csv("Table7_Prevalence_precision.csv"),
        "Table 7. Projected precision (%) at five attack prevalences, from equation (1).",
        "True-positive rate and false-alarm rate are held at the exact Table 6 fractions. The values are projections under that assumption.",
    )
    embed_table(
        doc, load_csv("Table8_Crossover_prevalence.csv"),
        "Table 8. Crossover attack prevalence from equation (3) at four miss-to-false-alarm cost ratios.",
        "Below the crossover the hybrid has lower expected cost than random forest; above it, random forest is the cheaper of those two learned detectors. The last two columns give the prevalences at which never alerting is cheaper than the hybrid and always alerting is cheaper than random forest. The ratios are scenarios, not measured prices. The inputs are the Table 6 counts, from unequal slices.",
    )
    embed_table(
        doc, load_csv("Table9_BoTIoT_seeds.csv"),
        "Table 9. BoT-IoT false alarms and misses on audited flow features in the 350,000-row mirror subsample. The first six count columns are the 0.5 cut. The last two are the seed-42 validation-tuned cut.",
        "Seed 42 has 72 benign test records. The tuned cut is the validation rule that allows one false alarm. For the seed-42 tree it leaves 175 misses. Tied scores leave logistic regression with two validation false alarms. The forest draws bags of 150,000 rows. One false alarm is about 1.4% of the benign test rows.",
    )
    embed_table(
        doc, load_csv("Table10_BoTIoT_features.csv"),
        "Table 10. BoT-IoT seed 42 on the same subsample: the 0.5 cut, the flow-only validation-tuned cut, and the same split with ports or IP-grouped aggregates restored.",
        "Addresses and timestamps stay out. The tuned row uses the threshold described in Table 9. IP aggregates are the published per-address fields, not the raw address strings. The tree assigns 0.965 of its importance to N_IN_Conn_P_DstIP when those fields are included.",
    )
    embed_table(
        doc, load_csv("Table11_CICDDoS_subset_ablation.csv"),
        "Table 11. Random forest on a four-file CICDDoS2019 subset, at the 0.5 cut, with medians fit on training rows only.",
        "The subset has 120,444 rows from UDPLag, Syn, NTP, and LDAP. It is not the matrix behind Table 3. The timestamp column is text and is already absent from the numeric features. Removing Inbound and the ports is the informative contrast. The hybrid was not refit on these rows.",
    )

    add_h(doc, "Figures")
    for stem in [
        "Figure_01_protocol",
        "Figure_02_reversal",
        "Figure_03_operating_points",
        "Figure_04_prevalence",
        "Figure_05_crossover",
        "Figure_06_botiot",
    ]:
        doc.add_picture(str(OUT_FIG / f"{stem}.png"), width=Inches(6.2))
        doc.add_paragraph()

    audit_callouts(doc)
    audit_citations(doc)
    doc.save(OUT / "02_Manuscript.docx")
    body = "\n".join(p.text for p in doc.paragraphs)
    return len(body.split()), len(ABSTRACT.split()), len(refs())


def audit_callouts(doc: Document):
    chunks = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "References":
            break
        chunks.append(paragraph.text)
    text = "\n".join(chunks)
    for kind, count in (("Table", 11), ("Figure", 6)):
        positions = []
        for number in range(1, count + 1):
            match = re.search(rf"\b{kind}\s+{number}\b", text)
            if not match:
                raise SystemExit(f"Missing in-text {kind} {number}")
            positions.append(match.start())
        if positions != sorted(positions):
            raise SystemExit(f"{kind} callouts out of order: {positions}")
    if re.search(r"\bTable\s+4\b", text.split("5. Results")[0]):
        raise SystemExit("Table 4 is cited before the results")
    banned = ["previous draft", "this revision", "not recovered", "2025 export folder", "do not read the results", "305"]
    for phrase in banned:
        if phrase.lower() in text.lower():
            raise SystemExit(f"Disallowed phrase still present: {phrase}")
    n_stored = len(re.findall(r"\bstored\b", text, flags=re.I))
    if n_stored > 2:
        raise SystemExit(f"'stored' appears {n_stored} times before the references")


def audit_citations(doc: Document):
    chunks = []
    for paragraph in doc.paragraphs:
        if paragraph.text.strip() == "References":
            break
        chunks.append(paragraph.text)
    main = "\n".join(chunks)
    # Family names that must appear. Dataset 'Canadian Institute' is cited as CIC.
    required = [
        "Akiba", "Anurupam", "Apruzzese", "Arp", "Axelsson", "Bergstra", "Bizzarri", "Chicco", "Cogranne",
        "Coscia", "D’hooge", "Doriguzzi-Corin", "Elkan", "ElShehaby", "Engelen", "Fawcett", "Ferrag", "Flood", "González-Ramos",
        "Grinsztajn", "Gu et al.", "Guo", "Hochreiter", "Hormozi", "Iqbal", "Javadpour", "Kapoor",
        "Ko et al.", "Koroniotis", "Lara-Gutierrez", "Layeghy", "Li et al.", "Madoune", "Meidan", "Mittal",
        "Neto", "Pedregosa", "Pendlebury", "Provost", "Ring", "Rustam", "Saito", "Sharafaldin",
        "Shi et al.", "Shwartz-Ziv", "Singha", "Snoek", "Sommer", "Vaarandi", "Wang et al.", "Zargar", "Zhao",
        "CIC, 2019",
    ]
    missing = [name for name in required if name not in main]
    if missing:
        raise SystemExit("Uncited references: " + ", ".join(missing))


def build_front_matter():
    doc = Document()
    add_para(doc, TITLE, size=14, bold=True, align="center", space_after=12)
    for name, mark, email, orcid, corresponding in AUTHORS:
        p = add_para(doc, align="center", space_after=2)
        run = p.add_run(name + ("*" if corresponding else ""))
        set_run_font(run, size=12, bold=True)
        sup = p.add_run(mark)
        set_run_font(sup, size=12, bold=True)
        sup.font.superscript = True
        add_para(doc, f"Email: {email}", size=10, align="center", space_after=2)
        add_para(doc, f"ORCID: {orcid}", size=10, align="center", space_after=6)
    p = add_para(doc, align="center", space_after=2)
    s = p.add_run("a")
    set_run_font(s, size=10)
    s.font.superscript = True
    t = p.add_run(" " + AFF["a"])
    set_run_font(t, size=10)
    p = add_para(doc, align="center", space_after=8)
    s = p.add_run("b")
    set_run_font(s, size=10)
    s.font.superscript = True
    t = p.add_run(" " + AFF["b"])
    set_run_font(t, size=10)
    add_para(doc, "* Corresponding author. Hina Iqbal, hinazoha4@gmail.com. " + AFF["a"], size=10, italic=True, align="center")
    add_para(doc, "Competing interests: The authors declare that they have no competing interests.", size=10)
    add_para(doc, "CRediT: H. Iqbal: Conceptualization; Methodology; Software; Formal analysis; Investigation; Data curation; Visualization; Writing – original draft; Project administration. C. Premebida: Supervision; Writing – review & editing; Validation. U. Haider: Writing – review & editing; Validation.", size=10)
    add_para(doc, "Funding: This research did not receive any specific grant from funding agencies in the public, commercial, or not-for-profit sectors.", size=10)
    doc.save(OUT / "01_Title_Page.docx")

    doc = Document()
    add_para(doc, "Highlights", size=14, bold=True, align="center")
    for item in HIGHLIGHTS:
        if len(item) > 85:
            raise SystemExit(f"Highlight too long ({len(item)}): {item}")
        add_para(doc, "• " + item, double=False, space_after=4)
    doc.save(OUT / "03_Highlights.docx")
    (OUT / "03_Highlights.txt").write_text("\n".join(HIGHLIGHTS) + "\n", encoding="utf-8")

    doc = Document()
    add_para(doc, "Cover letter", size=14, bold=True, align="center", space_after=10)
    add_para(doc, "17 October 2026", space_after=8)
    add_para(doc, "The Editors")
    add_para(doc, "Journal of Information Security and Applications", space_after=8)
    add_para(doc, "Dear Editors,", space_after=8)
    add_para(
        doc,
        f"Please consider our article “{TITLE}” for Journal of Information Security and Applications. The journal "
        "does not consider routine applications of existing machine-learning models to standard security datasets "
        "unless the work advances the security question itself. This manuscript takes that requirement as its "
        "scope. The models are a CNN–LSTM hybrid and standard tree ensembles. The security result is that their "
        "false-alarm order reverses between CICDDoS2019 and CICIoT2023 at the evaluated points, and that a prevalence-and-cost reading of "
        "the CICIoT2023 points changes which detector is cheaper. Later checks qualify that reading. On a four-file "
        "subset, removing an Inbound flag and the port fields raises the forest’s false alarms. On a BoT-IoT mirror "
        "subsample, a validation-tuned cut does not reduce the seed-42 tree’s misses. The analytical content applies "
        "established cost-sensitive decision theory to executed counts. It is an assessment of these detectors, not a new "
        "detection algorithm.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "On CICDDoS2019, counts recovered from the printed metrics give random forest no false alarms in 10,663 "
        "benign records and the hybrid 19 false alarms in 10,650 benign windows. Those binomial intervals do not "
        "overlap only if the windows are treated as independent, which overlapping windows violate. On CICIoT2023 "
        "the order reverses at one cut, and the two models were "
        f"scored on different cases: one hybrid false alarm in {NUMS['hybrid_benign']:,} benign windows, against "
        f"{NUMS['rf_fp']} random-forest false alarms in {NUMS['rf_benign']:,} benign records. Holding the CICIoT "
        f"rates fixed, projected precision at 1% attack prevalence is {NUMS['hybrid_precision_1pct']:.1f}% and "
        f"{NUMS['rf_precision_1pct']:.1f}%, and equal assumed costs switch the cheaper detector at "
        f"{NUMS['crossover_equal_pct']:.1f}% prevalence. The original experiments were carried out for the first author’s "
        "master’s dissertation at the University of Coimbra (Iqbal, 2026), which has not been published as a "
        "journal article. Section 5.5 replaces an earlier BoT-IoT table that used a truncated file. The new counts "
        "use a Kaggle mirror, keep all 477 benign rows, and subsample attacks to 350,000 rows. At seed 42 the tuned "
        "cut leaves the tree at 175 misses. Section 5.6 reports the four-file shortcut check. The original CIC "
        "comparison remains the originally scored slices. Derived materials for the historical calculations are at "
        "https://doi.org/10.5281/zenodo.23435601.",
        align="justify", first_line=True,
    )
    add_para(
        doc,
        "The manuscript is not under consideration elsewhere, and all authors have approved it. The authors declare "
        "no competing interests. The Zenodo record cited above is an interim deposit of the historical derived "
        "calculations. This submission’s Tables 9–11 supersede the BoT-IoT counts in that record. Raw datasets, "
        "training checkpoints, and the original CIC prediction scores are not included.",
        align="justify", first_line=True,
    )
    add_para(doc, "Sincerely,", space_after=8)
    add_para(doc, "Hina Iqbal")
    add_para(doc, "Corresponding author")
    add_para(doc, "hinazoha4@gmail.com")
    add_para(doc, "ORCID 0009-0002-2378-9414")
    add_para(doc, AFF["a"])
    add_para(doc, "Co-authors: Cristiano Premebida; Usman Haider")
    doc.save(OUT / "04_Cover_Letter.docx")

    statements = [
        ("05_Data_Availability.docx", "Data availability",
         "CICDDoS2019 and CICIoT2023 are available from the Canadian Institute for Cybersecurity, University of New Brunswick. "
         "The historical counts are those printed by the executed notebooks behind the first author’s master’s dissertation (Iqbal, 2026), available from the corresponding author on request. "
         "Derived tables and scripts for those calculations are deposited at https://doi.org/10.5281/zenodo.23435601. "
         "That interim record contains an earlier BoT-IoT run on a truncated file. Tables 9 and 10 in this manuscript replace it with a 350,000-row subsample of a complete Kaggle mirror that keeps all 477 benign rows. "
         "Table 11 is a four-file CICDDoS2019 subset. Raw dataset files are not redistributed. "
         "The original CIC score files, split indices, and model checkpoints are not in the deposit."),
        ("06_Competing_Interests.docx", "Declaration of competing interests",
         "The authors declare that they have no known competing financial interests or personal relationships that could have appeared to influence the work reported in this paper."),
        ("07_AI_Disclosure.docx", "Declaration of generative AI and AI-assisted technologies",
         "During the preparation of this work the authors used ChatGPT and Cursor for drafting and editing text, for restructuring sections, for assistance with plotting code, and for writing the BoT-IoT, CICIoT2023 pilot, and CICDDoS2019 subset scripts. "
         "The authors reviewed those scripts and executed them. Tables 9 and 10 come from the mirror subsample, Table 11 from the four-file subset, and Supplementary material S2 from the benign-versus-ICMP-flood pilot. "
         "Prevalence, crossover, and alert figures for the historical CIC results were computed from the confusion counts printed by the executed notebooks. "
         "The tools did not produce the original paired CICIoT2023 comparison or a hybrid refit on the reduced CICDDoS2019 features. "
         "The authors reviewed the manuscript and take full responsibility for the publication."),
    ]
    for filename, heading, text in statements:
        doc = Document()
        add_para(doc, heading, size=14, bold=True, align="center", space_after=10)
        add_para(doc, text, align="justify")
        doc.save(OUT / filename)

    bios = [
        ("Hina Iqbal",
         "Hina Iqbal is an MSc researcher in the Department of Electrical and Computer Engineering at the University of Coimbra, Portugal. Her work concerns evaluation practice for IoT intrusion detection, particularly false alarms, attack prevalence, and the operating points of DDoS detectors. Email: hinazoha4@gmail.com. ORCID: 0009-0002-2378-9414."),
        ("Cristiano Premebida",
         "Cristiano Premebida is with the Department of Electrical and Computer Engineering at the University of Coimbra, Portugal. His research interests include intelligent systems, machine learning, and perception for cyber-physical applications. He supervised the dissertation from which the detection rates are drawn. Email: cpremebida@deec.uc.pt. ORCID: 0000-0002-2168-2077."),
        ("Usman Haider",
         "Usman Haider is with the Department of Mechanics, Faculty of Civil Engineering, Czech Technical University in Prague, Czech Republic. His research interests include applied machine learning and data-driven modelling. He contributed to review and validation of the analyses in this article. Email: usmanhaider6886@gmail.com. ORCID: 0000-0002-0125-0912."),
    ]
    doc = Document()
    add_para(doc, "Author biographies", size=14, bold=True, align="center", space_after=10)
    for name, bio in bios:
        if len(bio.split()) > 100:
            raise SystemExit(f"Biography too long: {name}")
        add_para(doc, name, bold=True, space_after=2)
        add_para(doc, bio, align="justify", space_after=8)
    add_para(doc, "Passport-type photographs should be uploaded as separate files, as required by the journal.", size=10, italic=True)
    doc.save(OUT / "08_Author_Vitae.docx")


def write_zenodo(summary: dict):
    code = ZENODO / "code"
    for folder in (code, ZENODO / "tables", ZENODO / "figures", ZENODO / "predictions", ZENODO / "metadata"):
        folder.mkdir(parents=True, exist_ok=True)
    for name in ("derived_metrics.py", "reproduce_derived_analyses.py", "evaluate_saved_predictions.py", "run_botiot_audit.py"):
        shutil.copy(ROOT / "_build" / name, code / name)
    shutil.copy(ROOT / "_build" / "botiot_results.json", ZENODO / "metadata" / "botiot_results.json")
    shutil.copy(ROOT / "_build" / "cicddos_subset_report.json", ZENODO / "metadata" / "cicddos_subset_report.json")
    for csv_path in TABLES.glob("Table*.csv"):
        shutil.copy(csv_path, ZENODO / "tables" / csv_path.name)
    for fig in OUT_FIG.iterdir():
        if fig.suffix.lower() in {".png", ".pdf", ".tif"} and not fig.name.endswith("_proof.png"):
            shutil.copy(fig, ZENODO / "figures" / fig.name)
    (code / "requirements.txt").write_text(
        "numpy==2.5.1\nmatplotlib==3.11.0\npillow==12.3.0\n",
        encoding="utf-8",
    )
    (code / "README.md").write_text(
        "From this directory:\n\n"
        "    python reproduce_derived_analyses.py --tables ../tables --figures ../figures\n\n"
        "The script adds its own folder to the import path and uses Matplotlib’s DejaVu font. "
        "It rebuilds the prevalence, crossover, and alert figures from derived_metrics.py. "
        "Table 11 is read from ../metadata/cicddos_subset_report.json, and Tables 9 and 10 from "
        "../metadata/botiot_results.json. "
        "run_botiot_audit.py retrains the BoT-IoT models when the official CSVs are supplied locally. "
        "Those CSVs are not included. Delete any __pycache__ directory before a public upload.\n\n"
        "evaluate_saved_predictions.py computes miss rates at false-alarm budgets only when "
        "../predictions/scores.csv exists, with columns model, split, y_true, y_score and an optional seed.\n",
        encoding="utf-8",
    )
    (ZENODO / "predictions" / "README.md").write_text(
        "No score file is included.\n\n"
        "To compute matched false-alarm results, add scores.csv with columns model, split, y_true, y_score "
        "and an optional seed column, using split values val and test. "
        "Then run ../code/evaluate_saved_predictions.py.\n",
        encoding="utf-8",
    )
    (ZENODO / "LICENSE").write_text(
        "Creative Commons Attribution 4.0 International (CC BY 4.0)\n"
        "https://creativecommons.org/licenses/by/4.0/\n\n"
        "Derived tables, figures, and scripts in this deposit may be shared under CC BY 4.0.\n"
        "CICDDoS2019 and CICIoT2023 are not included and remain under the terms of the Canadian Institute for Cybersecurity.\n",
        encoding="utf-8",
    )
    (ZENODO / "CITATION.cff").write_text(
        "cff-version: 1.2.0\n"
        "message: Historical derived calculations are also deposited at https://doi.org/10.5281/zenodo.23435601. Tables 9-11 in this package supersede the BoT-IoT tables in that interim record.\n"
        f"title: \"Supporting material for: {TITLE}\"\n"
        "version: 1.0.0\n"
        "authors:\n"
        "  - family-names: Iqbal\n    given-names: Hina\n    orcid: https://orcid.org/0009-0002-2378-9414\n"
        "  - family-names: Premebida\n    given-names: Cristiano\n    orcid: https://orcid.org/0000-0002-2168-2077\n"
        "  - family-names: Haider\n    given-names: Usman\n    orcid: https://orcid.org/0000-0002-0125-0912\n"
        "date-released: 2026-10-17\n"
        "license: CC-BY-4.0\n"
        "type: software\n",
        encoding="utf-8",
    )
    (ZENODO / "metadata" / "environment.md").write_text(
        "# Environment used to generate the deposited figures\n\n"
        "Python packages, pinned to the versions that produced this deposit on 15 October 2026:\n\n"
        "- numpy==2.5.1\n- matplotlib==3.11.0\n- pillow==12.3.0\n\n"
        "Figures are written as PDF (vector), PNG, and LZW-compressed RGB TIFF at 300 dpi. "
        "The graphical abstract uses Matplotlib’s DejaVu Sans font, which ships with Matplotlib.\n\n"
        "The training notebooks recorded a Kaggle GPU session (Python 3.12). This deposit does not reconstruct that session and does not retrain the models.\n",
        encoding="utf-8",
    )
    (ZENODO / "metadata" / "source_and_licensing.md").write_text(
        "# Sources\n\n"
        "CICIoT2023 confusion counts and CICDDoS2019 comparison metrics: the executed notebooks behind "
        "Iqbal, H. (2026). Bayesian-optimized hybrid deep learning framework for detecting "
        "DDoS attacks on MEC and IoT networks. Master’s dissertation, University of Coimbra.\n\n"
        "CICDDoS2019 and CICIoT2023: Canadian Institute for Cybersecurity, University of New Brunswick. "
        "Obtain them from the Institute. They are not redistributed here.\n\n"
        "This deposit rebuilds the prevalence and cost analysis from the printed CICIoT2023 confusion counts "
        "and from CICDDoS2019 counts recovered from the printed precision, recall, support, and Matthews correlation. "
        "It does not contain raw traffic, model checkpoints, or prediction scores, and it does not retrain the models. "
        "The executed notebooks contain private download identifiers and are not included.\n\n"
        "BoT-IoT aggregate results are in metadata/botiot_results.json. They are a 350,000-row subsample of a Kaggle mirror, "
        "not the earlier truncated-file run. See metadata/botiot_source.md. Raw BoT-IoT files are not redistributed.\n"
        "The interim public deposit of the historical calculations is https://doi.org/10.5281/zenodo.23435601.\n",
        encoding="utf-8",
    )
    (ZENODO / "metadata" / "variable_dictionary.csv").write_text(
        "variable,meaning\n"
        "FAR,\"False-alarm rate, FP/(FP+TN)\"\n"
        "FNR,\"Miss rate, FN/(FN+TP)\"\n"
        "TPR,\"True-positive rate, equal to recall\"\n"
        "pi,Attack prevalence\n"
        "r,Miss cost divided by false-alarm cost\n"
        "pi_star,\"Crossover prevalence in equation (3)\"\n"
        "pi_never,\"Prevalence below which never alerting costs less than the hybrid\"\n"
        "pi_always,\"Prevalence above which always alerting costs less than random forest\"\n"
        "threshold,\"Reported decision: two-class argmax, equivalent to a softmax probability of one half except on ties\"\n",
        encoding="utf-8",
    )
    shutil.copy(ROOT / "05_Supplementary" / "S1_Literature_screen.md", ZENODO / "metadata" / "S1_Literature_screen.md")
    shutil.copy(ROOT / "05_Supplementary" / "S2_Protocol_pilot.md", ZENODO / "metadata" / "S2_Protocol_pilot.md")
    (ZENODO / "metadata" / "botiot_source.md").write_text(
        "# BoT-IoT source\n\n"
        "Official page: https://research.unsw.edu.au/projects/bot-iot-dataset\n\n"
        "Tables 9 and 10 are rebuilt from metadata/botiot_results.json. That file is a 350,000-row subsample "
        "of a Kaggle mirror of the 5% all-features extract (Koroniotis, Moustafa, Sitnikova and Turnbull, 2019). "
        "Every benign row is kept. The raw files are not included in this deposit.\n\n"
        "The mirror byte lengths differ from the published CloudStor lengths:\n\n"
        "| File | Published bytes | Mirror bytes |\n"
        "| --- | ---: | ---: |\n"
        "| UNSW_2018_IoT_Botnet_Full5pc_1.csv | 272335105 | 277543651 |\n"
        "| UNSW_2018_IoT_Botnet_Full5pc_2.csv | 273082424 | 278451723 |\n"
        "| UNSW_2018_IoT_Botnet_Full5pc_3.csv | 282131555 | 285316035 |\n"
        "| UNSW_2018_IoT_Botnet_Full5pc_4.csv | 189577740 | 191754639 |\n\n"
        "The mirror parses completely. It is not a byte-verified copy of the official files. "
        "An earlier truncated file of 268,435,456 bytes is not the source of Tables 9 and 10. "
        "Running code/run_botiot_audit.py on the full mirror will not reproduce those tables, "
        "because the deposited counts use the 350,000-row subsample.\n",
        encoding="utf-8",
    )
    meta = {
        "title": f"Supporting material for: {TITLE}",
        "version": "1.0.0",
        "upload_type": "software",
        "access_right": "open",
        "license": "cc-by-4.0",
        "creators": [
            {"name": "Iqbal, Hina", "affiliation": "University of Coimbra", "orcid": "0009-0002-2378-9414"},
            {"name": "Premebida, Cristiano", "affiliation": "University of Coimbra", "orcid": "0000-0002-2168-2077"},
            {"name": "Haider, Usman", "affiliation": "Czech Technical University in Prague", "orcid": "0000-0002-0125-0912"},
        ],
        "keywords": KEYWORDS,
        "description": (
            "Tables and scripts that rebuild a prevalence and expected-cost reading of executed DDoS detector counts, "
            "including the ROC convex hull and the never-alert and always-alert bounds. "
            "The scripts rebuild the derived tables, including the BoT-IoT subsample counts in metadata/botiot_results.json. "
            "The interim public deposit is https://doi.org/10.5281/zenodo.23435601. Licence: CC BY 4.0."
        ),
    }
    (ZENODO / "zenodo_metadata.json").write_text(json.dumps(meta, indent=2), encoding="utf-8")
    (ZENODO / "README.md").write_text(
        f"# Supporting material\n\n{TITLE}\n\n"
        "Version 1.0.0. This deposit rebuilds the paper’s prevalence, crossover, hull, and alert calculations "
        "from the confusion counts printed by the executed notebooks. It is not a training archive, and it does "
        "not reproduce the original model fits.\n\n"
        "## Contents\n\n"
        "- `tables/` — eleven tables\n"
        "- `figures/` — PDF, PNG, and LZW TIFF\n"
        "- `code/reproduce_derived_analyses.py` — rebuilds tables and figures from derived counts\n"
        "- `code/run_botiot_audit.py` — retrains the BoT-IoT comparison when the official CSVs are available locally\n"
        "- `code/evaluate_saved_predictions.py` — matched false-alarm evaluation when scores are supplied\n"
        "- `predictions/` — empty until a score file is added\n"
        "- `metadata/` — environment pins, BoT-IoT source note and results, literature screen\n\n"
        "## Reproduce\n\n"
        "From the deposit root, on Windows or Linux:\n\n"
        "```\n"
        "python code/reproduce_derived_analyses.py --tables tables_check --figures figures_check\n"
        "```\n\n"
        "Relative paths are enough. The script adds its own folder to the import path and draws figures with "
        "Matplotlib’s DejaVu Sans font.\n\n"
        "## Licence\n\n"
        "CC BY 4.0 for the derived tables, figures, and scripts. Raw CIC datasets are not included.\n\n"
        "## Authors\n\n"
        "Hina Iqbal, University of Coimbra, ORCID 0009-0002-2378-9414.\n"
        "Cristiano Premebida, University of Coimbra, ORCID 0000-0002-2168-2077.\n"
        "Usman Haider, Czech Technical University in Prague, ORCID 0000-0002-0125-0912.\n",
        encoding="utf-8",
    )
    return summary


def write_memo(words, abstract_words, n_refs, summary):
    (ROOT / "00_REVISION_MEMO.md").write_text(
        f"""# Internal memo

Internal note. Do not upload this file.

## Twentieth-revision package repair

No new model was trained. The full multi-attack CICIoT2023 comparison and the CICDDoS2019 hybrid refit remain blocked: the original split is not on this machine, and CUDA does not see the GeForce MX450.

- `cicddos_subset_report.json` is copied into the public Zenodo metadata folder. The reproduction script reads it from there, and from the folder beside the script during the build.
- Wolf et al. (2026), already in the JISA RIS file, is cited for the emulation-versus-hardware gap. Cogranne et al. (2026) and Singha et al. (2025) were already in the bibliography.
- The interim DOI https://doi.org/10.5281/zenodo.23435601 is still described as the earlier deposit. It is not described as containing Tables 9–11.
- A score of 95/100 is not claimed. Author photographs are still absent.

Abstract words: {abstract_words}. Paragraph words including references and captions: {words}. References: {n_refs}. RIS records screened: {summary['n']}.
""",
        encoding="utf-8",
    )
    (ROOT / "00_STATUS.txt").write_text(
        "Manuscript reports the repaired BoT-IoT subsample and the four-file CICDDoS subset. Not a new training run.\n"
        f"Abstract words: {abstract_words}. References: {n_refs}.\n"
        "The original paired CICIoT comparison and the hybrid refit were not run. Readiness is not 90/100.\n"
        "Upload the contents of 01_Submission_Files. Do not upload _build, this status file, or 00_REVISION_MEMO.md.\n"
        "Author photographs are still required and are not in this package.\n"
        "Kaggle output archives were not in the source folder.\n",
        encoding="utf-8",
    )


def sync():
    if SUB.exists():
        shutil.rmtree(SUB)
    shutil.copytree(OUT, SUB)
    sup = SUB / "05_Supplementary"
    if sup.exists():
        shutil.rmtree(sup)
    shutil.copytree(ROOT / "05_Supplementary", sup)
    for name in ("Graphical_Abstract.png", "Graphical_Abstract.tif"):
        shutil.copy(OUT_FIG / name, OUT / name)
        shutil.copy(OUT_FIG / name, SUB / name)
        duplicate = SUB / "Figures" / name
        if duplicate.exists():
            duplicate.unlink()
    for cache in ZENODO.rglob("__pycache__"):
        shutil.rmtree(cache)
    for compiled in ZENODO.rglob("*.pyc"):
        compiled.unlink()
    zip_path = ROOT / "20th_Revision_Submission.zip"
    if zip_path.exists():
        zip_path.unlink()
    shutil.make_archive(str(zip_path.with_suffix("")), "zip", root_dir=SUB)


def main():
    if len(ABSTRACT.split()) > 250:
        raise SystemExit(f"Abstract has {len(ABSTRACT.split())} words")
    for folder in (OUT, OUT_FIG, OUT_TAB, TABLES, ZENODO, ROOT / "05_Supplementary"):
        folder.mkdir(parents=True, exist_ok=True)
    write_all(TABLES, OUT_FIG)
    for csv_path in TABLES.glob("*.csv"):
        csv_to_xlsx(csv_path)
        shutil.copy(csv_path, OUT_TAB / csv_path.name)
        shutil.copy(csv_path.with_suffix(".xlsx"), OUT_TAB / (csv_path.stem + ".xlsx"))
    summary = screen_ris()
    if summary["n"] != 1691:
        raise SystemExit(f"Unexpected RIS count {summary['n']}")
    build_front_matter()
    words, abstract_words, n_refs = build_manuscript(summary)
    write_zenodo(summary)
    write_memo(words, abstract_words, n_refs, summary)
    sync()
    print("ABSTRACT", abstract_words, "WORDS", words, "REFS", n_refs)


if __name__ == "__main__":
    main()
